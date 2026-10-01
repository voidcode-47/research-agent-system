# -*- coding: utf-8 -*-
"""Markdown → Word 转换测试。"""
from io import BytesIO

from utils.md_to_docx import md_to_docx


class TestMdToDocx:
    def test_valid_docx_zip(self):
        buf = md_to_docx("# 标题\n\n正文内容")
        data = buf.getvalue()
        assert data[:2] == b"PK", "docx 应为 zip 容器"
        assert len(data) > 1500

    def test_roundtrip_structure(self):
        from docx import Document
        md = (
            "# 一级标题\n\n"
            "## 二级标题\n\n"
            "普通段落，包含**粗体**和[链接](https://doi.org/10.1000/abc)。\n\n"
            "- 列表项一\n"
            "- 列表项二\n\n"
            "> 引用内容\n\n"
            "---\n\n"
            "1. 有序项一\n"
            "2. 有序项二\n"
        )
        doc = Document(BytesIO(md_to_docx(md).getvalue()))
        texts = [p.text for p in doc.paragraphs]
        joined = "\n".join(texts)
        assert "一级标题" in joined
        assert "二级标题" in joined
        assert "粗体" in joined
        assert "https://doi.org/10.1000/abc" in joined  # 行内链接展开为 text (url)
        assert "列表项一" in joined
        assert "引用内容" in joined
        assert "有序项一" in joined
        # 标题级别检查
        headings = [p for p in doc.paragraphs if p.style.name.startswith("Heading")]
        assert headings and headings[0].style.name == "Heading 1"

    def test_empty_and_plain(self):
        buf = md_to_docx("")
        assert buf.getvalue()[:2] == b"PK"
        buf = md_to_docx("只有一行文字")
        assert "只有一行文字" in "\n".join(
            p.text for p in __import__("docx").Document(BytesIO(buf.getvalue())).paragraphs
        )
